import openpyxl,pathlib,json
root=pathlib.Path(__file__).resolve().parents[1]; out=root/'data/sheet_maps';out.mkdir(exist_ok=True)
for bank in (root/'documents').iterdir():
 paths=list(bank.glob('2026/2T26/*.xlsx'))
 for p in paths:
  if not any(s in p.name.lower() for s in ['serie','hist']):continue
  w=openpyxl.load_workbook(p,data_only=True,read_only=True); lines=[str(p.relative_to(root))]
  for s in w:
   lines.append('\nSHEET '+s.title)
   for row in s.iter_rows(max_row=350,max_col=min(s.max_column,180)):
    cells=[c for c in row if c.value is not None]
    if not cells:continue
    labels=[f'{c.coordinate}:{c.value}' for c in cells if isinstance(c.value,str)][:3]
    nums=[f'{c.coordinate}:{c.value} [{c.number_format}]' for c in cells if isinstance(c.value,(int,float))]
    lines.append(' | '.join(labels+nums[-2:]))
  (out/(bank.name+'.txt')).write_text('\n'.join(lines))
  print(bank.name,len(w.sheetnames),flush=True)
